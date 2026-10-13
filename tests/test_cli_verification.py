"""CLI: exit-code contract and truthful verification output.

Contract (README, dossier): 0 = every family ran and found no violation; 2 = violations (precedence); 3 = no violation but
verification incomplete (a family SKIPPED because its rule is off). Output always states the counts and outcome.
Callers of cli.main: the `slotsmith` console script only (pyproject.toml); no script or test depended on the old 0.
"""
from __future__ import annotations

import copy
from pathlib import Path

from slotsmith.cli import main
from slotsmith.rules import default_ruleset


def _files(tmp_path, r6_off=False):
    mdb = tmp_path / "mdb.xlsx"
    assert main(["generate", "--signals", "200", "--seed", "5", "--errors", "3", "--out", str(mdb)]) == 0
    rules = None
    if r6_off:
        r = copy.deepcopy(default_ruleset())
        next(t for t in r.toggles if t.id == "R6-hardware").enabled = False
        rules = tmp_path / "r6_off.json"
        rules.write_text(r.model_dump_json())
    return mdb, rules


def test_allocate_full_pass_exits_0_without_incomplete_warning(tmp_path, capsys):
    mdb, _ = _files(tmp_path)
    out = tmp_path / "a.xlsx"
    assert main(["allocate", str(mdb), "--out", str(out)]) == 0
    assert "incomplete" not in capsys.readouterr().err
    assert main(["verify", str(out), "--mdb", str(mdb)]) == 0
    assert "-> passed" in capsys.readouterr().out


def test_incomplete_verification_is_reported_not_hidden(tmp_path, capsys):
    mdb, rules = _files(tmp_path, r6_off=True)
    out = tmp_path / "a.xlsx"
    code = main(["allocate", str(mdb), "--rules", str(rules), "--out", str(out)])
    err = capsys.readouterr().err
    assert code == 3                                    # incomplete verification is not success
    assert "-> incomplete" in err and "1 skipped" in err and "H-1..6" in err
    assert main(["verify", str(out), "--mdb", str(mdb), "--rules", str(rules)]) == 3
    text = capsys.readouterr().out
    assert "SKIPPED" in text and "not run: rule R6-hardware is switched off" in text and "-> incomplete" in text


def test_violation_exits_2(tmp_path, capsys):
    mdb, _ = _files(tmp_path)
    out = tmp_path / "a.xlsx"
    assert main(["allocate", str(mdb), "--out", str(out)]) == 0
    capsys.readouterr()
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb["IO Allocation"]
    hdr = [c.value for c in ws[1]]
    ti = hdr.index("Terminal Block")
    for row in ws.iter_rows(min_row=2):
        if row[ti].value == "Fused-injected-power":
            row[ti].value = "Standard"
            break
    wb.save(out)
    assert main(["verify", str(out), "--mdb", str(mdb)]) == 2
    assert "-> failed" in capsys.readouterr().out


def test_revise_with_skipped_family_exits_3(tmp_path, capsys):
    mdb, rules = _files(tmp_path, r6_off=True)
    revb = tmp_path / "revb.xlsx"
    assert main(["generate", "--signals", "200", "--seed", "5", "--errors", "3", "--out", str(tmp_path / "x.xlsx"),
                 "--rev-b", str(revb)]) == 0
    prev = tmp_path / "a.xlsx"
    assert main(["allocate", str(mdb), "--rules", str(rules), "--out", str(prev)]) == 3
    capsys.readouterr()
    code = main(["revise", "--prev", str(prev), "--mdb", str(revb), "--rules", str(rules), "--out", str(tmp_path / "b.xlsx")])
    assert code == 3 and "-> incomplete" in capsys.readouterr().err


def test_violation_takes_precedence_over_incomplete(tmp_path, capsys):
    from slotsmith.cli import _exit_code
    chk = [{"Rule": "K-1", "Result": "FAIL", "Violations": 2}, {"Rule": "H-1..6", "Result": "SKIPPED", "Violations": 0}]
    assert _exit_code(chk) == 2
    assert _exit_code([{"Rule": "K-1", "Result": "PASS", "Violations": 0}]) == 0
    assert _exit_code([{"Rule": "K-1", "Result": "PASS", "Violations": 0}, {"Rule": "H", "Result": "SKIPPED", "Violations": 0}]) == 3
