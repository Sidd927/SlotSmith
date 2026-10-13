"""R6-hardware and verification truthfulness.

Documented (docs/SlotSmith_Technical_Dossier.md, rule register section 6): an I/O type with no CHARM/terminal-block
mapping raises the blocking `H-*` technical query. NOT documented anywhere: what switching R6 off means for such a row.

Chosen interpretation (fail closed, decided 13 Oct 2026, not organiser-approved): the `H-*` query and the held row are
unconditional. R6 only controls whether the verifier compares CHARM/terminal block with the table; when off, that family is
reported SKIPPED (never PASS) and the run is an "incomplete" verification. No hardware is ever invented.
"""
from __future__ import annotations

import copy
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from slotsmith.generator import generate, make_revision
from slotsmith.pipeline import run
from slotsmith.revision import revise
from slotsmith.rules import default_ruleset
from slotsmith.verifier import SKIPPED, summarize, verify

BAD_TAG_NOTE = "PULSE-XYZ"


def _workbook() -> Path:
    g = generate(300, seed=7, inject_errors=6)     # one injected row has the unmapped IO type 'PULSE-XYZ'
    p = Path(tempfile.mkdtemp()) / "mdb.xlsx"
    g.df.to_excel(p, index=False)
    return p


def _rules(**off):
    r = copy.deepcopy(default_ruleset())
    for t in r.toggles:
        if t.id in off.get("switches", ()):
            t.enabled = False
    return r


def _run(rules, out=True):
    p = Path(tempfile.mkdtemp()) / "out.xlsx" if out else None
    return run([_workbook()], rules, p), p


def _unresolved(res):
    hw = [q for q in res.tqs if q.rule_ref == "H-*"]
    assert hw, "the unmapped IO type must raise an H-* technical query"
    rows = [r for r in res.rows if r["TQ Ref"] in {q.id for q in hw}]
    assert rows
    return hw, rows


@pytest.fixture(scope="module")
def on():
    return _run(_rules())


@pytest.fixture(scope="module")
def off():
    return _run(_rules(switches=("R6-hardware",)))


# 1/2. unsupported IO type, rule enabled and disabled: same fail-closed outcome
@pytest.mark.parametrize("which", ["on", "off"])
def test_unmapped_io_type_is_held_never_allocated(which, request):
    res, _ = request.getfixturevalue(which)
    hw, rows = _unresolved(res)
    assert all(q.severity == "blocking" and q.category == "hardware" for q in hw)
    for r in rows:
        assert r["Status"] == "TQ raised" and r["Slot"] == "" and r["CHARM Type"] == "" and r["Terminal Block"] == ""
        assert r["Tag"] and r["Source"] and r["IO Type"] == BAD_TAG_NOTE         # 3. traceable, original row preserved
    assert not any(r["Status"] == "Allocated" and not r["CHARM Type"] for r in res.rows)   # nothing allocated without hardware


def test_disabled_message_says_unknown_hardware_stays_unresolved(off):
    hw, _ = _unresolved(off[0])
    assert "R6-hardware is switched off" in hw[0].message and "unresolved" in hw[0].message


# 5/6. verification states
def test_disabled_family_is_skipped_not_pass(off):
    res, _ = off
    h = next(c for c in res.check if c["Rule"] == "H-1..6")
    assert h["Result"] == SKIPPED and h["Violations"] == 0 and "R6-hardware" in h["Note"]
    sm = summarize(res.check)
    assert (sm["passed"], sm["failed"], sm["skipped"], sm["total"]) == (10, 0, 1, 11)
    assert sm["outcome"] == "incomplete"          # never "passed" while a family was skipped


def test_all_enabled_run_has_no_skips_and_passes(on):
    sm = summarize(on[0].check)
    assert (sm["passed"], sm["failed"], sm["skipped"]) == (11, 0, 0) and sm["outcome"] == "passed"


@pytest.mark.parametrize("sw,rid", [("R2-segregation", "S-1"), ("R1-grouping", "G-1"), ("R5-spare", "P-1")])
def test_other_switchable_families_are_skipped_when_off(sw, rid):
    res, _ = _run(_rules(switches=(sw,)), out=False)
    c = next(c for c in res.check if c["Rule"] == rid)
    assert c["Result"] == SKIPPED and sw in c["Note"]
    sm = summarize(res.check)
    assert sm["skipped"] == 1 and sm["outcome"] == "incomplete"      # never "passed"; M-1 now passes (see tests/test_mixed_cable.py)


def test_group_isolation_none_skips_g1():
    r = _rules()
    r.group_isolation_level = "none"
    res, _ = _run(r, out=False)
    assert next(c for c in res.check if c["Rule"] == "G-1")["Result"] == SKIPPED


def test_app_never_certifies_a_run_with_skipped_checks(off, on):
    from app.core.state import is_verifier_certified, verification_status
    ok, msg, failed = is_verifier_certified(off[0])
    assert ok is False and "skipped" in msg and "Not a full verification" in msg and failed == []
    assert verification_status(off[0])["outcome"] == "incomplete"
    ok_on, msg_on, _ = is_verifier_certified(on[0])
    assert ok_on is True and "all 11 rule families ran and passed" in msg_on


# 4/7. exported evidence
def test_workbook_records_tq_switches_and_verification(off):
    res, path = off
    wb = openpyxl.load_workbook(path, read_only=True)
    tq = [r for r in wb["TQ List"].iter_rows(values_only=True)]
    assert any(r[3] == "H-*" and BAD_TAG_NOTE in str(r[4]) for r in tq)
    sw = {r[0]: r[1] for r in wb["Rule Set"].iter_rows(min_row=2, values_only=True) if r and str(r[0]).startswith("rule switch")}
    assert sw["rule switch R6-hardware"] == "OFF" and sw["rule switch R2-segregation"] == "ON"
    info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(min_row=2, values_only=True)}
    assert info["disabled rule switches"] == "R6-hardware"
    assert "10 passed, 0 failed, 1 skipped" in info["verification"] and "outcome: incomplete" in info["verification"]
    assert "H-1..6 not run" in info["verification"] and info["unresolved hardware (H-* queries)"].startswith("1 held")
    rc = [r for r in wb["Rule Check"].iter_rows(values_only=True)]
    h = next(r for r in rc if r[0] == "H-1..6")
    assert h[2] == "SKIPPED" and "R6-hardware" in str(h[5])
    io = [r for r in wb["IO Allocation"].iter_rows(values_only=True)]
    hdr = list(io[0])
    row = next(r for r in io[1:] if r[hdr.index("IO Type")] == BAD_TAG_NOTE)
    assert row[hdr.index("Status")] == "TQ raised" and row[hdr.index("TQ Ref")] and row[hdr.index("CHARM Type")] is None


def test_enabled_workbook_reports_no_disabled_rules_and_full_pass(on):
    wb = openpyxl.load_workbook(on[1], read_only=True)
    info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(min_row=2, values_only=True)}
    assert info["disabled rule switches"] == "none" and "11 passed, 0 failed, 0 skipped" in info["verification"]


# 8. strict re-verification
def _plain(rows):
    return [{k: (None if v in ("", None) else str(v)) for k, v in r.items()} for r in rows]


def test_strict_reverification_of_disabled_run_runs_h_and_passes_held_rows(off):
    res, _ = off
    strict = verify(_plain(res.rows), res.mdb.df, default_ruleset())
    sm = summarize(strict)
    assert sm["skipped"] == 0 and sm["outcome"] == "passed"      # held row has a TQ; nothing is claimed as placed


def test_strict_verifier_flags_hardware_blanked_on_a_placed_row(off):
    res, _ = off
    rows = _plain(res.rows)
    victim = next(r for r in rows if r["Row Kind"] == "Signal" and r["Status"] == "Allocated")
    victim["CHARM Type"] = victim["Terminal Block"] = None
    strict = verify(rows, res.mdb.df, default_ruleset())
    assert next(c for c in strict if c["Rule"] == "H-1..6")["Result"] == "FAIL"
    # even with R6 off, a placed signal without hardware is a Q-1 completeness violation (fail closed)
    loose = verify(rows, res.mdb.df, _rules(switches=("R6-hardware",)))
    q1 = next(c for c in loose if c["Rule"] == "Q-1")
    assert q1["Result"] == "FAIL" and "unresolved hardware" in q1["Examples"]
    assert next(c for c in loose if c["Rule"] == "H-1..6")["Result"] == SKIPPED


# 9. revision keeps the unresolved state
@pytest.mark.parametrize("sw", [(), ("R6-hardware",)])
def test_revision_keeps_unresolved_rows_unresolved(sw):
    from tests.test_core import _canon
    rules = _rules(switches=sw)
    g = generate(300, seed=7, inject_errors=6)
    a = run(None, rules, df=_canon(g.df), verify_from_file=False)
    rr = revise(a.rows, _canon(make_revision(g.df, seed=8)), rules)
    hw = [q for q in rr.tqs if q.rule_ref == "H-*"]
    assert hw and all(q.severity == "blocking" for q in hw)
    assert not any(r["Status"] == "Allocated" and not r["CHARM Type"] for r in rr.rows)
    assert any(r["Status"] == "TQ raised" and r["IO Type"] == BAD_TAG_NOTE for r in rr.rows)
    assert summarize(rr.check)["skipped"] == len(sw)


# 11/12. normal mappings and enabled-rule rejection still work (also covered by tests/test_core.py)
def test_normal_hardware_mappings_unchanged(on):
    res, _ = on
    by_io = {}
    for r in res.rows:
        if r["Row Kind"] == "Signal" and r["Status"] == "Allocated":
            by_io.setdefault(r["IO Type"], set()).add((r["CHARM Type"], r["Terminal Block"]))
    rules = default_ruleset()
    for io, got in by_io.items():
        row = rules.hardware_for(rules.normalise_io_type(io))
        assert got == {(row.charm, row.terminal_block)}


def test_mutated_terminal_block_still_fails_when_r6_on(on):
    res, _ = on
    rows = _plain(res.rows)
    r = next(r for r in rows if r["Terminal Block"] == "Fused-injected-power")
    r["Terminal Block"] = "Standard"
    assert next(c for c in verify(rows, res.mdb.df, default_ruleset()) if c["Rule"] == "H-1..6")["Result"] == "FAIL"


def test_build_rows_has_no_unchecked_path_to_blank_hardware():
    import inspect
    from slotsmith import export
    src = inspect.getsource(export.build_rows)
    assert "reached export without a hardware rule" in src     # invariant raises instead of emitting a blank allocated row


# ----------------------------------------------------------------------------- end-to-end consistency (real workbooks)
def test_ui_session_workbooks_agree_with_in_memory_status():
    """Rev A and Rev B workbooks written by the app session carry the same verification story as the objects the UI reads."""
    from app.core.session import ensure_session, load_synthetic, run_allocation, run_revision, set_revision_synthetic
    from app.core.state import unresolved_hardware, verification_status
    d: dict = {}
    ensure_session(d)
    load_synthetic(d, 300, 3, 6)
    rules = copy.deepcopy(d["rules"])
    next(t for t in rules.toggles if t.id == "R6-hardware").enabled = False
    d["rules"] = rules
    assert run_allocation(d)[0] and set_revision_synthetic(d) and run_revision(d)[0]
    for path, obj in ((d["out"], d["res"]), (d["rev_out"], d["rr"])):
        vs = verification_status(obj)
        wb = openpyxl.load_workbook(path, read_only=True)
        info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(min_row=2, values_only=True)}
        assert f"{vs['passed']} passed, {vs['failed']} failed, {vs['skipped']} skipped" in info["verification"]
        assert f"outcome: {vs['outcome']}" in info["verification"] and info["disabled rule switches"] == "R6-hardware"
        hw = unresolved_hardware(obj)
        assert info["unresolved hardware (H-* queries)"].startswith(f"{len(hw)} held")
        rc = {r[0]: r[2] for r in wb["Rule Check"].iter_rows(min_row=2, values_only=True)}
        assert rc["H-1..6"] == "SKIPPED" and sum(v == "SKIPPED" for v in rc.values()) == vs["skipped"] == 1
        tq_rules = [r[3] for r in wb["TQ List"].iter_rows(min_row=2, values_only=True)]
        assert tq_rules.count("H-*") == len(hw) >= 1
