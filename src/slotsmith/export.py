"""Turn an allocation into flat rows and write the Excel deliverable.

The organiser's IO allocation template has not been released (docs/known_limitations.md),
so the workbook layout here follows the PS output description and is driven by
COLUMNS, which can be re-mapped to the real template when it arrives.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .allocator import AllocationResult
from .rules import RuleSet
from .validate import ValidationResult

COLUMNS = [
    "Tag", "IO Type", "Cable", "Cable Size", "Pair", "IO Group", "Segregation", "Row Kind",
    "Controller", "System Cabinet", "Marshalling Cabinet", "CIOC", "Baseplate", "Slot",
    "CHARM Address", "CHARM Type", "Terminal Block", "Status", "TQ Ref", "Source",
]


def controller_map(result: AllocationResult, rules: RuleSet) -> dict[tuple[int, int], tuple[int, int, int]]:
    """(cabinet, cioc_pos) -> (global CIOC seq, controller seq, system cabinet seq)."""
    out = {}
    seq = 0
    for ci, cab in enumerate(result.layout.cabinets, 1):
        for qi, cioc in enumerate(cab.ciocs, 1):
            seq += 1
            no = cioc.no if cioc.no is not None else seq
            ctl = cioc.ctl if cioc.ctl is not None else (no - 1) // rules.max_ciocs_per_controller + 1
            syscab = (ctl - 1) // rules.controllers_per_system_cabinet + 1
            out[(ci, qi)] = (no, ctl, syscab)
    return out


def build_rows(val: ValidationResult, result: AllocationResult, rules: RuleSet) -> list[dict]:
    n = rules.naming
    fmt = lambda pat, **kw: pat.format(area=n.area, unit=n.unit, **kw)
    cmap = controller_map(result, rules)
    cap = rules.charms_per_baseplate
    rows: list[dict] = []
    blocks = {b.cable: b for b in val.blocks}
    for cable, pl in result.placements.items():
        b = blocks[cable]
        seq, ctl, syscab = cmap[(pl.cabinet, pl.cioc)]
        for pair in range(1, b.size + 1):
            lin = (pl.start_slot - 1) + (pair - 1)
            bp = pl.baseplate + lin // cap
            slot = lin % cap + 1
            s = b.signals.get(pair)
            if s is not None:
                hw = rules.hardware_for(s.io_type)
                kind, tag, io, status, tq, src = "Signal", s.tag, s.io_type_raw, "Allocated", "", s.src
                if hw is None:   # unreachable: validate() holds every signal without a hardware rule (fail closed)
                    raise ValueError(f"{s.tag}: IO type '{s.io_type_raw}' reached export without a hardware rule")
                charm, tb = hw.charm, hw.terminal_block
            else:
                held = pair in val.held_slots.get(cable, set())
                kind = "Held (TQ)" if held else "Spare pair"
                tag, io, status, tq, src = "", "", "Reserved" if held else "Spare", "", ""
                charm, tb = "", ""
            rows.append({
                "Tag": tag, "IO Type": io, "Cable": cable, "Cable Size": b.size, "Pair": pair,
                "IO Group": (s.group if s is not None and s.group else b.group) or "",
                "Segregation": (s.seg if s is not None and s.seg else b.seg) or "", "Row Kind": kind,
                "Controller": fmt(n.controller, seq=ctl),
                "System Cabinet": fmt(n.system_cabinet, seq=syscab),
                "Marshalling Cabinet": fmt(n.marshalling_cabinet, seq=pl.cabinet),
                "CIOC": fmt(n.cioc, seq=seq),
                "Baseplate": fmt(n.baseplate, cioc_no=seq, n=bp),
                "Slot": slot,
                "CHARM Address": fmt(n.charm_slot, cioc_no=seq, n=bp, slot=slot),
                "CHARM Type": charm, "Terminal Block": tb, "Status": status, "TQ Ref": tq, "Source": src,
                "_cab": pl.cabinet, "_cioc": pl.cioc, "_bp": bp,
            })
    # Signals that could not be placed: one row each, with the TQ that explains why.
    placed_tags = {r["Tag"] for r in rows if r["Tag"]}
    for b in val.blocks:
        if b.cable in result.unplaced:
            for s in b.signals.values():
                rows.append(_unplaced_row(s.tag, s.io_type_raw, b.cable, b.size, s.pair, b.group, b.seg,
                                          "Not allocated", "capacity", s.src))
    return rows


def _unplaced_row(tag, io, cable, size, pair, group, seg, status, tq, src):
    return {"Tag": tag, "IO Type": io, "Cable": cable, "Cable Size": size, "Pair": pair, "IO Group": group or "",
            "Segregation": seg or "", "Row Kind": "Signal", "Controller": "", "System Cabinet": "",
            "Marshalling Cabinet": "", "CIOC": "", "Baseplate": "", "Slot": "", "CHARM Address": "",
            "CHARM Type": "", "Terminal Block": "", "Status": status, "TQ Ref": tq, "Source": src,
            "_cab": None, "_cioc": None, "_bp": None}


def _raw(v):
    """Original cell value, with pandas NaN/None shown as empty (a NaN is truthy and would otherwise be written as 'nan')."""
    return "" if v is None or (isinstance(v, float) and v != v) else v


def add_blocked_rows(rows: list[dict], df: pd.DataFrame, val: ValidationResult) -> list[dict]:
    """Every MDB row whose signal is blocked by a TQ gets an explicit 'TQ raised' row (rule Q-1).

    The lookup key must match validate(): a blank tag (NaN from a CSV/Excel read) is keyed by its source row."""
    from .validate import _clean
    for r in df.itertuples(index=False):
        src = f"{r.src_file}:{r.src_sheet}:{r.src_row}"
        key = _clean(r.tag) or src
        if key in val.blocked_tags:
            rows.append(_unplaced_row(_raw(r.tag), _raw(r.io_type), _raw(r.cable), "", _raw(r.pair),
                                      _raw(r.group), _raw(r.seg), "TQ raised", val.blocked_tags[key], src))
    return rows


def summary(rows: list[dict], result: AllocationResult, rules: RuleSet) -> list[dict]:
    cap_slots = rules.baseplates_per_cabinet * rules.charms_per_baseplate
    by_cab = defaultdict(lambda: {"sig": 0, "spare": 0, "held": 0, "bps": set(), "ciocs": set(), "name": ""})
    for r in rows:
        if r["_cab"] is None:
            continue
        d = by_cab[r["_cab"]]
        d["name"] = r["Marshalling Cabinet"]
        d["bps"].add((r["_cioc"], r["_bp"]))
        d["ciocs"].add(r["_cioc"])
        d["sig" if r["Row Kind"] == "Signal" else "spare" if r["Row Kind"] == "Spare pair" else "held"] += 1
    out = []
    for cab in sorted(by_cab):
        d = by_cab[cab]
        used_bp = len(d["bps"])
        free_slots = used_bp * rules.charms_per_baseplate - d["sig"] - d["spare"] - d["held"]
        out.append({
            "Marshalling Cabinet": d["name"], "CIOCs": len(d["ciocs"]), "Baseplates used": used_bp,
            "Empty baseplates (cabinet spare)": rules.baseplates_per_cabinet - used_bp,
            "Signals": d["sig"], "Wired spare pairs": d["spare"], "Held for TQ": d["held"],
            "Free slots on used baseplates": free_slots,
            "Fill % (signals / cabinet slots)": round(100 * d["sig"] / cap_slots, 1),
            "Fill % (occupied / cabinet slots)": round(100 * (d["sig"] + d["spare"] + d["held"]) / cap_slots, 1),
        })
    return out


def _write_sheet(ws, records: list[dict], cols: list[str] | None = None):
    cols = cols or (list(records[0].keys()) if records else ["(none)"])
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5F2A")
        c.alignment = Alignment(vertical="center")
    for rec in records:
        ws.append([rec.get(c, "") for c in cols])
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(rec.get(c, ""))) for rec in records[:500]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _evidence(tqs, check: list[dict], rules: RuleSet) -> dict:
    """Run Info lines that make the workbook self-explanatory: effective switches, verification counts, unresolved hardware."""
    from .verifier import summarize
    off = [t.id for t in rules.toggles if not t.enabled]
    out = {"disabled rule switches": ", ".join(off) if off else "none"}
    if check:
        sm = summarize(check)
        txt = f"{sm['total']} families: {sm['passed']} passed, {sm['failed']} failed, {sm['skipped']} skipped; outcome: {sm['outcome']}"
        if sm["skipped_rules"]:
            txt += " (skipped: " + "; ".join(f"{r} {n}" for r, n in sm["skipped_rules"]) + ")"
        out["verification"] = txt
    else:
        out["verification"] = "not yet run when this sheet was first written"
    hw = [t for t in tqs if t.rule_ref == "H-*"]
    out["unresolved hardware (H-* queries)"] = (f"{len(hw)} held: " + ", ".join(sorted({x for t in hw for x in t.tags}))[:300]
                                                 if hw else "none")
    out["technical queries"] = f"{len(tqs)} total, {sum(1 for t in tqs if t.severity == 'blocking')} blocking"
    return out


def write_workbook(path: str | Path, rows: list[dict], summ: list[dict], tqs, check: list[dict],
                   rules: RuleSet, meta: dict, impact: list[dict] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "IO Allocation"
    key = lambda r: (r["_cab"] or 9999, r["_cioc"] or 0, r["_bp"] or 0, r["Slot"] or 0, r["Tag"])
    _write_sheet(ws, sorted(rows, key=key), COLUMNS)
    _write_sheet(wb.create_sheet("Cabinet Summary"), summ)
    _write_sheet(wb.create_sheet("TQ List"), [
        {"TQ": t.id, "Severity": t.severity, "Category": t.category, "Rule": t.rule_ref, "Message": t.message,
         "Tags": ", ".join(t.tags[:20]), "Cables": ", ".join(t.cables[:20]), "Source rows": ", ".join(t.sources[:20]),
         "Suggested action": t.suggested_action} for t in tqs] or [{"TQ": "none"}])
    _write_sheet(wb.create_sheet("Rule Check"), check)
    if impact is not None:
        _write_sheet(wb.create_sheet("Revision Impact"), impact or [{"Change": "none"}])
    switches = [{"Parameter": f"rule switch {t.id}", "Value": "ON" if t.enabled else "OFF",
                 } for t in rules.toggles]
    _write_sheet(wb.create_sheet("Rule Set"), [{"Parameter": k, "Value": str(v)} for k, v in
                                               rules.model_dump(exclude={"hardware", "toggles", "io_type_aliases"}).items()] +
                 switches +
                 [{"Parameter": f"hardware {h.id}", "Value": f"{h.io_types} -> {h.charm} / {h.terminal_block}"} for h in rules.hardware])
    _write_sheet(wb.create_sheet("Run Info"), [{"Key": k, "Value": str(v)} for k, v in {**meta, **_evidence(tqs, check, rules)}.items()])
    wb.save(path)
