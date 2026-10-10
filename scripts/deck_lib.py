"""Small helpers on top of python-pptx for the SlotSmith Round 1 deck.

The official template is published only as a PDF, so each slide uses the
rendered template page as a full-bleed background (preserving branding and the
mandatory headings exactly) and all content is added as editable shapes/text.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

SLIDE_W = 13.333
SLIDE_H = 13.333 * 703.5 / 1257.75  # template PDF aspect ratio (≈7.458 in)

# Palette sampled from the template (accent green bar = #4DD262)
GREEN = "4DD262"
GREEN_DK = "1F5F2A"
PANEL = "06140A"
PANEL2 = "0C2412"
TEXT = "EAF5EC"
MUTED = "A9C2AE"
AMBER = "FFB547"
RED = "FF6B6B"
CYAN = "6FD3FF"
WHITE = "FFFFFF"

BODY_FONT = "Arial"

# IO-type colours used in every cabinet/baseplate graphic
IO_COLORS = {
    "AI": "4FA3FF",
    "AO": "B98CFF",
    "DI": "7BD88F",
    "DO": "FFD54F",
    "DO_ISO": "FF8A50",
    "DO_RELAY": "FF6FA8",
    "SPARE": "5B6B5F",
    "FREE": "0E1A12",
}


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def add_bg_slide(prs, bg_path: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    pic = slide.shapes.add_picture(bg_path, 0, 0, prs.slide_width, prs.slide_height)
    pic.name = "Template background"
    return slide


def _set_alpha(fill_elm_parent, alpha_pct: float):
    """Apply transparency to the solidFill colour under a shape's spPr."""
    srgb = fill_elm_parent.find(".//" + qn("a:srgbClr"))
    if srgb is None:
        return
    for a in srgb.findall(qn("a:alpha")):
        srgb.remove(a)
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int((100 - alpha_pct) * 1000)))


def box(slide, x, y, w, h, fill=PANEL, line=None, line_w=1.0, transparency=0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None, dash=False):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if transparency:
            _set_alpha(shp._element.spPr, transparency)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            ln = shp.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if name:
        shp.name = name
    return shp


def text(slide, x, y, w, h, paras, size=13, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=BODY_FONT, spacing_after=2, line_spacing=1.0,
         margin=0.0, name=None, shape=None):
    """paras: str | list[str | list[run]] where run = (text, {opts}).

    opts: size, color, bold, italic, font.
    A paragraph may be a dict {"runs": [...], "bullet": True, "level": 0, "align": ...}.
    """
    if shape is None:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        tb = shape
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p_align = align
        bullet = False
        runs = para
        if isinstance(para, dict):
            runs = para["runs"]
            bullet = para.get("bullet", False)
            p_align = para.get("align", align)
        if isinstance(runs, str):
            runs = [(runs, {})]
        p.alignment = p_align
        p.space_after = Pt(spacing_after)
        p.line_spacing = line_spacing
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.16))))
            pPr.set("indent", str(-Emu(Inches(0.16))))
            bu = etree.SubElement(pPr, qn("a:buFont"))
            bu.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar"))
            bc.set("char", "•")
        for rt, opts in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = rgb(opts.get("color", color))
    if name:
        tb.name = name
    return tb


def arrow(slide, x1, y1, x2, y2, color=GREEN, width=1.5, dash=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    if name:
        c.name = name
    return c


def notes(slide, txt: str):
    slide.notes_slide.notes_text_frame.text = txt
